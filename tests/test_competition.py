import pandas as pd
import pytest
from openpyxl import load_workbook

from competition.build_dataset import build_dataset
from competition.sanitize import FORBIDDEN_PUBLIC_COLS, check_public_columns
from voltvision.assumptions import deep_merge, load_template
from voltvision.simulate import build_panel, replay_claims

from conftest import FROZEN_TEMPLATE

ROUNDS = [
    {'name': 'base', 'label': 'Base world', 'patch': {}},
    {'name': 'freq_p10', 'label': 'Frequency +10%', 'patch': {'frequency_intensity': 1.10}},
    {'name': 'severity_shock', 'label': 'Severity +50%', 'patch': {'severity_multiplier': 1.50}},
]


@pytest.fixture(scope='module')
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp('dataset')
    public, truth, master = build_dataset(
        ROUNDS, load_template(FROZEN_TEMPLATE), out,
        train_years=2, test_years=1, observed_seed=42, scoring_seeds=[67, 69],
        panel_seed=20260916, n=60, verbose=False)
    return public, truth, master


def _cfg():
    base = load_template(FROZEN_TEMPLATE)
    return deep_merge(base, {'n': 60, 'n_years': 3, 'seed': 42})


def test_replay_is_deterministic_and_attributes_fixed():
    cfg = _cfg()
    panel = build_panel(cfg, 20260916, 3)
    b1, c1 = replay_claims(panel, cfg, 67)
    b2, c2 = replay_claims(panel, cfg, 67)
    assert b1.equals(b2) and c1.equals(c2)

    b3, _ = replay_claims(panel, cfg, 69)
    keys = ['POLID', 'SIM_YEAR']
    a = b1.set_index(keys)[['SUM_ASSURED', 'DRIVER_AGE', 'CAR_AGE']]
    c = b3.set_index(keys)[['SUM_ASSURED', 'DRIVER_AGE', 'CAR_AGE']]
    assert a.equals(c)
    assert not b1['CLAIM_AMOUNT'].equals(b3['CLAIM_AMOUNT'])


def test_public_workbook_has_no_leaks(built):
    public, _, _ = built
    wb = load_workbook(public, read_only=True)
    assert {'Instructions', 'Train_Policies', 'Train_Claims', 'Test_Policies',
            'Submission', 'Scenarios'} <= set(wb.sheetnames)

    test = pd.read_excel(public, sheet_name='Test_Policies')
    assert not FORBIDDEN_PUBLIC_COLS.intersection(test.columns)
    assert not [c for c in test.columns if c.startswith('CLAIM_')]
    assert test['TEST_ID'].is_unique
    assert list(dict.fromkeys(test['SCENARIO'])) == ['base', 'freq_p10', 'severity_shock']

    train_p = pd.read_excel(public, sheet_name='Train_Policies')
    train_c = pd.read_excel(public, sheet_name='Train_Claims')
    assert not FORBIDDEN_PUBLIC_COLS.intersection(train_p.columns)
    assert not FORBIDDEN_PUBLIC_COLS.intersection(train_c.columns)
    assert {'RAW_GAMMA', 'CAP'}.isdisjoint(train_c.columns)

    sub = pd.read_excel(public, sheet_name='Submission', skiprows=2)
    assert list(sub.columns[:2]) == ['TEST_ID', 'PREMIUM']
    assert len(sub) == len(test)


def test_truth_covers_every_scenario(built):
    public, truth_path, _ = built
    truth = pd.read_excel(truth_path, sheet_name='Truth')
    test = pd.read_excel(public, sheet_name='Test_Policies')
    assert set(truth['TEST_ID']) == set(test['TEST_ID'])
    assert [c for c in truth.columns if c.startswith('CLAIMS_S')] == ['CLAIMS_S1', 'CLAIMS_S2']
    assert set(truth['SCENARIO']) == {'base', 'freq_p10', 'severity_shock'}
    assert (truth[['CLAIMS_S1', 'CLAIMS_S2']] >= 0).all().all()


def test_master_workbook_formulas(built):
    public, truth_path, master = built
    wb = load_workbook(master)
    assert {'Master sheet', 'Dataset', 'Train_Policies', 'Train_Claims', 'Truth',
            'Team A', 'Team B', 'Team T', 'Subs', 'Alloc', 'By_Scenario',
            'Checks', 'Instructions'} <= set(wb.sheetnames)

    ds = wb['Dataset']
    truth = pd.read_excel(truth_path, sheet_name='Truth')
    assert ds.max_row - 1 == len(truth)

    ms = wb['Master sheet']
    assert ms['A3'].value == 'Scenario: base'
    assert ms['C3'].value == 'Team A'
    assert ms['C4'].value.startswith('=IF(By_Scenario!$C2=0')
    assert ms['C5'].value == '=By_Scenario!$C2'
    assert ms['C6'].value == '=By_Scenario!$F2'
    assert ms['C7'].value == '=IF(By_Scenario!$G2="OK",TRUE,FALSE)'
    assert ms['C8'].value.startswith('=IF(C7=FALSE')
    assert ms['A24'].value == 'OVERALL'

    assert wb['Subs']['A2'].value == '=Dataset!B2'
    assert wb['Subs']['B2'].value == "=IF('Team A'!B2<=0,\"\",'Team A'!B2)"
    assert wb['Team A']['A2'].value == '=Dataset!B2'
    assert wb['Alloc']['B2'].value.startswith('=IF(Subs!B2=')
    assert wb['By_Scenario']['C2'].value.startswith('=SUMPRODUCT(--(Subs!$W$2')
    assert ',Alloc!B$2:B$' in wb['By_Scenario']['C2'].value
    assert ',Subs!B$2:B$' in wb['By_Scenario']['D2'].value
    assert ',Subs!$V$2:$V$' in wb['By_Scenario']['E2'].value
    assert 'N(' not in wb['By_Scenario']['D2'].value
    assert '*' not in wb['By_Scenario']['D2'].value
    assert '$A:$A' not in wb['Subs']['V2'].value
    assert 'Truth!$A$2:$A$' in wb['Subs']['V2'].value
    assert wb['By_Scenario']['G2'].value == '=IF(D2=0,"no wins",IF(E2>=D2,"DQ","OK"))'
    assert str(wb['Checks']['B1'].value).startswith('=COUNTA(Dataset!')

    assert len(wb['Team A'].data_validations.dataValidation) == 1
    assert wb['By_Scenario'].max_row == 1 + 3 * 20 + 20


def test_check_public_columns_rejects_leaks():
    with pytest.raises(ValueError):
        check_public_columns(['POLID', 'CLAIM_LAMBDA'], 'ctx')


def test_team_count_and_names_are_flexible(tmp_path):
    out = tmp_path / 'custom'
    _, _, master = build_dataset(
        ROUNDS, load_template(FROZEN_TEMPLATE), out,
        train_years=2, test_years=1, observed_seed=42, scoring_seeds=[67],
        panel_seed=20260916, n=40, team_names=['Alpha', 'Beta', 'Gamma'],
        verbose=False)
    wb = load_workbook(master)
    assert {'Alpha', 'Beta', 'Gamma'} <= set(wb.sheetnames)
    assert 'Team A' not in wb.sheetnames
    assert wb['Master sheet']['C3'].value == 'Alpha'
    assert wb['Master sheet']['E3'].value == 'Gamma'
    assert wb['Master sheet']['F3'].value is None
    assert wb['Subs']['B2'].value == "=IF('Alpha'!B2<=0,\"\",'Alpha'!B2)"
    assert wb['Subs']['D2'].value == "=IF('Gamma'!B2<=0,\"\",'Gamma'!B2)"
    assert wb['By_Scenario'].max_row == 1 + 3 * 3 + 3


def test_team_sheet_name_sanitized():
    from competition.workbooks import _sheet_names, _sheet_ref
    assert _sheet_names(['A/B', 'A?B']) == ['A_B', 'A_B~01']
    assert _sheet_ref("O'Brien") == "'O''Brien'"
